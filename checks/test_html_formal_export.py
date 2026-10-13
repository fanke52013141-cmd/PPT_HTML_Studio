"""S05: real browser/media tests, isolated fixture approval; no model calls."""

import json
from pathlib import Path
from types import SimpleNamespace
import shutil
import subprocess

import pytest

from html_build_snapshot import SnapshotError, fresh_hash
from html_export_snapshot import freeze_course, verify_course, page_folder
from html_render_runner import (
    HtmlRenderRunner,
    HtmlRenderRunnerDependencies,
    HtmlRenderError,
)
from html_visual_store import save_scene
from html_visual_review_service import approve_scene, HtmlReviewDependencies
from runtime_support import run_subprocess_killable
from tts_artifacts import build_confirmation_payload

ROOT = Path(__file__).resolve().parents[1]


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


@pytest.fixture
def course(tmp_path):
    run = tmp_path / "run"
    ids = ["s1", "s2", "s3"]
    source = json.loads(
        (ROOT / "html_engine/visual/scenes/condensation.json").read_text(
            encoding="utf-8"
        )
    )
    for index, slide in enumerate(ids):
        scene = {**source, "id": slide}
        save_scene(run, slide, scene, index)
        write(
            run / f"planning/html_visual/binding-{slide}.json",
            {
                "format": "hps.html.motion_binding",
                "version": "0.1.0",
                "mode": "beat_ids",
                "actions": {},
            },
        )
        directory = run / "slides" / slide
        directory.mkdir(parents=True)
        subprocess.run(
            [
                "ffmpeg",
                "-v",
                "error",
                "-y",
                "-f",
                "lavfi",
                "-i",
                "anullsrc=r=48000:cl=stereo",
                "-t",
                "1.2",
                str(directory / "voice.mp3"),
            ],
            check=True,
        )
        (directory / "tts_text.txt").write_text("观察变化。", encoding="utf-8")
        (directory / "subtitles.srt").write_text(
            "1\n00:00:00,000 --> 00:00:01,100\n观察变化。\n", encoding="utf-8"
        )
        write(directory / "tts_metadata.json", {"fixture": True})
        write(
            directory / "audio_timeline.json",
            {
                "audio_content_duration_sec": 1.2,
                "segments": [
                    {"id": "b1", "start": 0, "end": 1.1, "text": "观察变化。"}
                ],
            },
        )
        approve_scene(
            scene,
            run_dir=run,
            deps=HtmlReviewDependencies(repo_root=ROOT, json_generator=lambda **kw: {}),
            reviewer="test_fixture",
        )
    write(
        run / "planning/visual_contract.json",
        {
            "slides": [
                {
                    "slide_id": slide,
                    "narration_beats": [{"id": "b1", "spoken_text": "观察变化。"}],
                }
                for slide in ids
            ]
        },
    )
    write(
        run / "planning/audio_confirmed.json",
        build_confirmation_payload(run, ids, confirmation_mode="test_fixture"),
    )
    return run, ids


def runner(execute=run_subprocess_killable):
    return HtmlRenderRunner(
        HtmlRenderRunnerDependencies(
            ROOT, lambda run: [], run_subprocess_bounded=execute
        )
    )


def cleanup(result):
    shutil.rmtree(result.export_evidence["work_dir"])


def test_pptx_freeze_order_and_tamper(course):
    run, ids = course
    frozen = freeze_course(
        run, ROOT, target="pptx", fingerprint_factory=lambda: {"digest": "test"}
    )
    original = verify_course(frozen)
    assert [p["slide_id"] for p in original["pages"]] == ids
    page = page_folder(frozen, original["pages"][0])
    scene = json.loads((page / "input/scene.json").read_text(encoding="utf-8"))
    assert scene["beats"] == []
    write(run / "planning/visual_contract.json", {"slides": [{"slide_id": ids[-1]}]})
    assert verify_course(frozen)["pages"] == original["pages"]
    (page / "preview/player.js").write_text("tampered")
    with pytest.raises(SnapshotError, match="HASH"):
        verify_course(frozen)


def test_third_page_failure_resume_reorder_and_corrupt_cache(course):
    run, ids = course
    fail = True
    rendered = []

    def execute(cmd, **kwargs):
        nonlocal fail
        if len(cmd) > 1 and str(cmd[1]).endswith("export-slide-video.cjs"):
            slide = json.loads(Path(cmd[2]).read_text(encoding="utf-8"))["id"]
            rendered.append(slide)
            if slide == "s3" and fail:
                return subprocess.CompletedProcess(cmd, 1, "", "injected page3 failure")
        return run_subprocess_killable(cmd, **kwargs)

    production = runner(execute)
    project = SimpleNamespace(run_dir=str(run))
    with pytest.raises(HtmlRenderError, match="page3"):
        production.run(
            project, output_dir=run / "videos", set_stage=lambda stage: None, fps=2
        )
    assert len(list((run / "planning/html_visual/segment-cache").glob("*.mp4"))) == 2
    assert not list((run / "planning/html_visual/export-work").iterdir())
    fail = False
    rendered.clear()
    result = production.run(
        project, output_dir=run / "videos", set_stage=lambda stage: None, fps=2
    )
    assert rendered == ["s3"]
    assert [p["reused"] for p in result.export_evidence["pages"]] == [True, True, False]
    assert result.export_evidence["validation"]["full_decode"] == "passed"
    assert [p["offset_frames"] for p in result.export_evidence["pages"]] == [0, 3, 6]
    assert not (run / "videos").exists()  # candidate remains private staging
    cache_path = (
        run
        / "planning/html_visual/segment-cache"
        / (result.export_evidence["pages"][0]["cache_key"] + ".mp4")
    )
    cleanup(result)
    contract = json.loads(
        (run / "planning/visual_contract.json").read_text(encoding="utf-8")
    )
    contract["slides"] = contract["slides"][:2][::-1]
    write(run / "planning/visual_contract.json", contract)
    write(
        run / "planning/audio_confirmed.json",
        build_confirmation_payload(
            run, ids[:2][::-1], confirmation_mode="test_fixture"
        ),
    )
    rendered.clear()
    reordered = production.run(
        project, output_dir=run / "videos", set_stage=lambda stage: None, fps=2
    )
    assert not rendered
    assert [p["slide_id"] for p in reordered.export_evidence["pages"]] == ["s2", "s1"]
    assert reordered.export_evidence["validation"]["frames"] == 6
    cleanup(reordered)
    cache_path.write_bytes(b"corrupt")
    recovered = production.run(
        project, output_dir=run / "videos", set_stage=lambda stage: None, fps=2
    )
    assert rendered == ["s1"]
    assert recovered.export_evidence["pages"][1]["reused"] is False
    cleanup(recovered)


def test_complete_decode_rejects_failure(tmp_path):
    from html_frozen_video import validate_video

    fake = runner(
        lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, "", "corrupt midpoint")
    )
    fake._probe = lambda *args: {
        "streams": [
            {"codec_type": "video", "nb_frames": "30", "codec_name": "h264"},
            {"codec_type": "audio", "codec_name": "aac", "duration": "1"},
        ],
        "format": {"duration": "1"},
    }
    with pytest.raises(ValueError, match="DECODE"):
        validate_video(fake, tmp_path / "candidate.mp4", 30, 30)


def test_parallel_runs_have_independent_candidates_and_cache_lock(course):
    from concurrent.futures import ThreadPoolExecutor

    run, ids = course
    contract = json.loads(
        (run / "planning/visual_contract.json").read_text(encoding="utf-8")
    )
    contract["slides"] = contract["slides"][:1]
    write(run / "planning/visual_contract.json", contract)
    write(
        run / "planning/audio_confirmed.json",
        build_confirmation_payload(run, ids[:1], confirmation_mode="test_fixture"),
    )

    def render():
        return runner().run(
            SimpleNamespace(run_dir=str(run)),
            output_dir=run / "videos",
            set_stage=lambda stage: None,
            fps=2,
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: render(), range(2)))
    assert results[0].output_path != results[1].output_path
    assert sorted(r.export_evidence["pages"][0]["reused"] for r in results) == [
        False,
        True,
    ]
    assert all(r.export_evidence["validation"]["frames"] == 3 for r in results)
    cleanup(results[0])
    assert results[1].output_path.is_file()
    cleanup(results[1])


def test_pptx_uses_frozen_text_when_live_author_changes(course):
    from artifact_fingerprint import presentation_input_fingerprint
    from html_pptx_snapshot import export_frozen_pptx, HtmlSnapshotDependencies
    from pptx import Presentation
    import hashlib

    run, ids = course
    before = presentation_input_fingerprint(run)
    changed = False

    def execute(cmd, **kwargs):
        nonlocal changed
        if str(cmd[1]).endswith("export-snapshots.cjs") and not changed:
            frozen = json.loads(Path(cmd[2]).read_text(encoding="utf-8"))
            author_path = run / "planning/html_visual/scene-s1.json"
            author = json.loads(author_path.read_text(encoding="utf-8"))
            title = next(n for n in author["nodes"] if n["type"] == "text")
            title["runs"][0]["text"] = "新版本文字"
            write(author_path, author)
            assert frozen["nodes"] != author["nodes"]
            assert "HPS_HTML_RESOURCES" not in kwargs["env"]
            changed = True
        return run_subprocess_killable(cmd, **kwargs)

    exported = export_frozen_pptx(
        run,
        ROOT,
        "frozen.pptx",
        lambda: presentation_input_fingerprint(run),
        deps=HtmlSnapshotDependencies(ROOT, run_subprocess_bounded=execute),
    )
    assert changed and exported["fingerprint"] == before
    assert exported["fingerprint"] != presentation_input_fingerprint(run)
    deck = Presentation(exported["path"])
    for slide, page in zip(
        deck.slides, exported["metadata"]["export_evidence"]["pages"]
    ):
        assert (
            hashlib.sha256(slide.shapes[0].image.blob).hexdigest()
            == page["image_sha256"]
        )
    assert not list((run / "planning/html_visual/export-work").iterdir())


def test_default_fps_and_real_frozen_presenter_composite(course, monkeypatch):
    from test_video_task_status_consistency import _build_service, _RecordingJobStore
    from html_frozen_video import validate_video

    run, ids = course
    contract = json.loads(
        (run / "planning/visual_contract.json").read_text(encoding="utf-8")
    )
    contract["slides"] = contract["slides"][:1]
    write(run / "planning/visual_contract.json", contract)
    write(
        run / "planning/audio_confirmed.json",
        build_confirmation_payload(run, ids[:1], confirmation_mode="test_fixture"),
    )
    source = run / "planning/digital_human/digi_upload.mp4"
    source.parent.mkdir(parents=True)
    subprocess.run(
        [
            "ffmpeg",
            "-v",
            "error",
            "-y",
            "-f",
            "lavfi",
            "-i",
            "color=c=gray:s=160x160:r=30",
            "-t",
            "1.5",
            "-pix_fmt",
            "yuv420p",
            str(source),
        ],
        check=True,
    )
    write(run / "planning/digital_human.json", {"enabled": True, "mode": "upload"})
    production = runner()
    result = production.run(
        SimpleNamespace(run_dir=str(run)),
        output_dir=run / "videos",
        set_stage=lambda stage: None,
    )
    assert result.export_evidence["validation"]["fps"] == 30
    import math
    # MP3 encoder padding varies by installed FFmpeg. Check the actual input
    # duration, rather than assuming that every build emits a 37th frame.
    audio_probe = json.loads(subprocess.run(
        ["ffprobe", "-v", "error", "-show_format", "-of", "json", str(run / "slides/s1/voice.mp3")],
        capture_output=True, check=True,
    ).stdout)
    expected_frames = math.ceil(max(1.2, float(audio_probe["format"]["duration"])) * 30)
    assert result.export_evidence["validation"]["frames"] == expected_frames
    frozen = Path(result.export_evidence["frozen_run_dir"])
    before = fresh_hash(source)
    source.write_bytes(b"changed live source")
    write(run / "planning/digital_human.json", {"enabled": False})
    assert fresh_hash(frozen / "planning/digital_human/digi_upload.mp4") == before
    service, _, _, _ = _build_service(
        run, run / "videos/final.mp4", _RecordingJobStore()
    )
    object.__setattr__(service.dependencies, "html_runner", production)
    service._set_task_stage = lambda *args: None
    service.runner._validate_render_color = lambda *args, **kwargs: pytest.fail(
        "HTML must not use Remotion color gates"
    )
    composed, used = service._apply_digital_human_composite(
        SimpleNamespace(id="fixture", run_dir=str(frozen)), result, "fixture-task"
    )
    assert used and composed.color_validation["full_decode"] == "passed"
    _, validation = validate_video(production, composed.output_path, 30, expected_frames)
    assert validation["full_decode"] == "passed"
    cleanup(result)
